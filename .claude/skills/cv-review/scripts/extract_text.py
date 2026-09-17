#!/usr/bin/env python
"""Estrae il testo di uno o piu' CV (PDF, DOCX, MD, TXT, HTML) in Markdown/testo.

Uso:
    python extract_text.py <file-o-cartella> [...] [--out DIR]

- Un file: stampa il testo su stdout.
- Piu' file o una cartella: scrive <nome>.txt in --out (default: ./cv_text/).
- Le immagini nei PDF vengono ignorate; se il PDF e' una scansione senza layer
  di testo il risultato e' vuoto e viene segnalato su stderr.

Dipendenze (una qualsiasi di queste per i PDF): pdftotext (poppler), pymupdf, pypdf.
Per i DOCX: pandoc oppure python-docx.
"""
import argparse
import shutil
import subprocess
import sys
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".doc", ".md", ".txt", ".html", ".htm", ".rtf", ".odt"}


def _pdf(path: Path) -> str:
    if shutil.which("pdftotext"):
        r = subprocess.run(["pdftotext", "-layout", str(path), "-"], capture_output=True, text=True)
        if r.returncode == 0 and r.stdout.strip():
            return r.stdout
    try:
        import fitz  # pymupdf
        with fitz.open(path) as doc:
            return "\n".join(page.get_text() for page in doc)
    except ImportError:
        pass
    try:
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)
    except ImportError:
        pass
    raise RuntimeError("nessun estrattore PDF disponibile (installa poppler, pymupdf o pypdf)")


def _pandoc(path: Path) -> str:
    if not shutil.which("pandoc"):
        raise RuntimeError("pandoc non disponibile")
    r = subprocess.run(["pandoc", str(path), "-t", "gfm", "--wrap=none"], capture_output=True, text=True)
    if r.returncode != 0:
        raise RuntimeError(r.stderr.strip())
    return r.stdout


def _docx(path: Path) -> str:
    try:
        return _pandoc(path)
    except RuntimeError:
        pass
    try:
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)
    except ImportError:
        raise RuntimeError("ne' pandoc ne' python-docx disponibili per i DOCX")


def extract(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        text = _pdf(path)
    elif ext in {".docx", ".doc", ".odt", ".rtf", ".html", ".htm"}:
        text = _docx(path) if ext == ".docx" else _pandoc(path)
    else:
        text = path.read_text(encoding="utf-8", errors="replace")
    if not text.strip():
        print(f"[avviso] {path.name}: nessun testo estratto (PDF scansionato? serve OCR)", file=sys.stderr)
    return text


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("inputs", nargs="+", help="file o cartelle")
    ap.add_argument("--out", default="cv_text", help="cartella di output per piu' file")
    args = ap.parse_args()

    files = []
    for inp in args.inputs:
        p = Path(inp)
        if p.is_dir():
            files += sorted(f for f in p.iterdir() if f.suffix.lower() in SUPPORTED)
        elif p.exists():
            files.append(p)
        else:
            print(f"[errore] non trovato: {p}", file=sys.stderr)

    if not files:
        print("[errore] nessun file da elaborare", file=sys.stderr)
        return 1

    if len(files) == 1:
        sys.stdout.write(extract(files[0]))
        return 0

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    for f in files:
        try:
            target = out / f"{f.name}.txt"
            target.write_text(extract(f), encoding="utf-8")
            print(f"ok  {f.name} -> {target}")
        except Exception as e:  # noqa: BLE001
            print(f"ERR {f.name}: {e}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
